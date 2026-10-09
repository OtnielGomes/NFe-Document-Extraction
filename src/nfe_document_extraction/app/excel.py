"""Workbook export of the document the demo already holds. No Gradio."""

from __future__ import annotations

import io
import tempfile
from collections.abc import Mapping
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.worksheet.worksheet import Worksheet

from nfe_document_extraction.app.copy import (
    field_label,
    line_headers,
    outcome_label,
    status_label,
    visitor_text,
)

_HEADER_MONEY = frozenset(
    {
        "total_products",
        "freight",
        "insurance",
        "discount",
        "ipi_total",
        "other_expenses",
        "icms_total",
        "icms_st_total",
        "invoice_total",
    }
)
_FLAGGED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
_LINE_FIELDS = (
    "product_code",
    "description",
    "ncm",
    "cfop",
    "quantity",
    "unit_price",
    "item_total",
)
_LINE_TEXT = frozenset({"product_code", "description", "ncm", "cfop"})


def excel_filename(body: Mapping[str, object]) -> str:
    """``nfe-<number>-<series>.xlsx``, or ``nfe-extraction.xlsx`` when either is unread."""
    number = _header_identifier(body, "number")
    series = _header_identifier(body, "series")
    if number is None or series is None:
        return "nfe-extraction.xlsx"
    return f"nfe-{number}-{series}.xlsx"


def excel_download_path(body: Mapping[str, object], *, language: str = "pt") -> Path:
    """Write the workbook where Gradio can serve it. The file name is the download name."""
    directory = Path(tempfile.mkdtemp(prefix="nfe-excel-"))
    path = directory / excel_filename(body)
    path.write_bytes(workbook_bytes(body, language=language))
    return path


def workbook_bytes(body: Mapping[str, object], *, language: str = "pt") -> bytes:
    """Three sheets: summary, header fields, and line items in printed order."""
    book = Workbook()
    summary = book.active
    if summary is None:
        raise RuntimeError("workbook has no active sheet")
    summary.title = visitor_text("Summary", language)
    _write_summary(summary, body, language)
    _write_header(book.create_sheet(visitor_text("Header", language)), body, language)
    _write_line_items(
        book.create_sheet(visitor_text("Line items", language)),
        body,
        language,
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_summary(sheet: Worksheet, body: Mapping[str, object], language: str) -> None:
    outcome = body.get("outcome")
    sheet["A1"] = visitor_text("Outcome", language)
    sheet["B1"] = outcome_label(str(outcome), language)
    sheet["A2"] = visitor_text("Review messages", language)
    messages = body.get("review_messages")
    if not isinstance(messages, list):
        return
    for offset, message in enumerate(messages):
        sheet.cell(
            row=2 + offset,
            column=2,
            value=visitor_text(str(message), language),
        )


def _write_header(sheet: Worksheet, body: Mapping[str, object], language: str) -> None:
    sheet.append(
        [
            visitor_text("Field", language),
            visitor_text("Value", language),
            visitor_text("Status", language),
            visitor_text("Messages", language),
        ]
    )
    extraction = _extraction(body)
    for field in _fields(body):
        path = _path(field)
        if not path.startswith("header."):
            continue
        value, number_format = _header_cell(extraction, path)
        sheet.append(
            [
                field_label(path, language, extraction),
                value,
                status_label(_status(field), language),
                _messages(field, language),
            ]
        )
        sheet.cell(row=sheet.max_row, column=2).number_format = number_format


def _write_line_items(
    sheet: Worksheet, body: Mapping[str, object], language: str
) -> None:
    sheet.append(list(line_headers(language)))
    items = _extraction(body).get("line_items")
    if not isinstance(items, list):
        raise TypeError("Line items must be a list.")
    flagged = {
        path
        for field in _fields(body)
        if _status(field) == "flagged"
        for path in (_path(field),)
    }
    for index, item in enumerate(items):
        if not isinstance(item, dict):
            raise TypeError("A line item must be an object.")
        sheet.append([None] * len(_LINE_FIELDS))
        row_index = sheet.max_row
        for column, name in enumerate(_LINE_FIELDS, start=1):
            cell = sheet.cell(row=row_index, column=column)
            cell.value, cell.number_format = _line_cell(item, name)
            if f"line_items[{index}].{name}" in flagged:
                cell.fill = _FLAGGED_FILL


def _header_cell(extraction: Mapping[str, object], path: str) -> tuple[object, str]:
    name = path.removeprefix("header.")
    raw = _nested(extraction, path)
    if name == "issue_date":
        return _date(raw), "YYYY-MM-DD"
    if name in _HEADER_MONEY:
        return _reais(raw), "0.00"
    return _text(raw), FORMAT_TEXT


def _line_cell(item: Mapping[str, object], name: str) -> tuple[object, str]:
    raw = item.get(name)
    if name == "item_total":
        return _reais(raw), "0.00"
    if name not in _LINE_TEXT:
        amount = _decimal(raw)
        return float(amount), _decimal_format(amount)
    return _text(raw), FORMAT_TEXT


def _header_identifier(body: Mapping[str, object], name: str) -> str | None:
    extraction = body.get("extraction")
    if not isinstance(extraction, dict):
        return None
    header = extraction.get("header")
    if not isinstance(header, dict):
        return None
    value = header.get(name)
    if isinstance(value, str) and value.isdigit():
        return value
    return None


def _extraction(body: Mapping[str, object]) -> Mapping[str, object]:
    extraction = body.get("extraction")
    if not isinstance(extraction, dict):
        raise TypeError("A document needs an extraction.")
    return extraction


def _fields(body: Mapping[str, object]) -> tuple[Mapping[str, object], ...]:
    fields = body.get("fields")
    if not isinstance(fields, list):
        raise TypeError("A document needs a field list.")
    rows: list[Mapping[str, object]] = []
    for field in fields:
        if not isinstance(field, dict):
            raise TypeError("A field entry must be an object.")
        rows.append(field)
    return tuple(rows)


def _path(field: Mapping[str, object]) -> str:
    path = field.get("field")
    if not isinstance(path, str):
        raise TypeError("A field path must be a string.")
    return path


def _status(field: Mapping[str, object]) -> str:
    status = field.get("status")
    if not isinstance(status, str):
        raise TypeError("A field status must be a string.")
    return status


def _messages(field: Mapping[str, object], language: str) -> str:
    messages = field.get("messages")
    if not isinstance(messages, list):
        raise TypeError("Field messages must be a list.")
    return "; ".join(visitor_text(str(message), language) for message in messages)


def _nested(extraction: Mapping[str, object], path: str) -> object:
    current: object = extraction
    for part in path.split("."):
        if not isinstance(current, dict) or part not in current:
            raise KeyError(path)
        current = current[part]
    return current


def _reais(value: object) -> float:
    if isinstance(value, bool) or not isinstance(value, int):
        raise TypeError("Money must be integer cents.")
    return float(Decimal(value) / Decimal(100))


def _decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, str):
        return Decimal(value)
    raise TypeError("Quantity and unit price must be decimals.")


def _decimal_format(amount: Decimal) -> str:
    exponent = amount.as_tuple().exponent
    if not isinstance(exponent, int) or exponent >= 0:
        return "0"
    return "0." + ("0" * -exponent)


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date.fromisoformat(value[:10])
    raise TypeError("Issue date must be a date.")


def _text(value: object) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if value is None:
        return ""
    return str(value)
