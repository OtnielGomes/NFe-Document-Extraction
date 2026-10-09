"""Presentation copy: Portuguese by default, English selectable, sources stay English."""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import gradio as gr
import pytest
from openpyxl import load_workbook
from test_api import RecordingMonitor, ScriptedReader, _pdf
from test_identifier_rules import ACCESS_KEY, _extraction, _line_item
from test_ui import _OVER_CAP_OUTPUT_TOKENS, _service

from nfe_document_extraction.agents.sections import SectionReading, SectionRequest
from nfe_document_extraction.app.copy import (
    HEIC_RULE,
    UPLOAD_RULE,
    field_label,
    gallery_caption,
    intro_body,
    visitor_text,
)
from nfe_document_extraction.app.errors import RequestRejected
from nfe_document_extraction.app.excel import excel_filename, workbook_bytes
from nfe_document_extraction.app.gallery import load_gallery
from nfe_document_extraction.app.limits import CostLimits
from nfe_document_extraction.app.ui import (
    DemoActions,
    ScreenCopy,
    build_blocks,
    field_rows,
    language_chrome,
    present_upload,
    present_visitor_feedback,
    render_result,
    screen_copy,
)
from nfe_document_extraction.app.uploads import (
    HEIC_NOT_ACCEPTED,
    UNSUPPORTED_TYPE,
    prepare_upload,
)
from nfe_document_extraction.evaluation.metrics import ExtractionFailure
from nfe_document_extraction.validation import validate

_TODAY = date(2026, 10, 6)
_IP = "198.51.100.20"
_HEIC = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00"
_SYNTHETIC = Path("data/synthetic")

_MESSAGES = {
    "Issuer CNPJ must contain 14 digits.": "O CNPJ do emitente deve ter 14 dígitos.",
    "Issuer CNPJ cannot be a sequence of repeated digits.": (
        "O CNPJ do emitente não pode ser uma sequência de dígitos repetidos."
    ),
    "Issuer CNPJ check digit does not match.": (
        "O dígito verificador do CNPJ do emitente não confere."
    ),
    "Recipient CNPJ/CPF must contain 11 digits (CPF) or 14 digits (CNPJ).": (
        "O CNPJ/CPF do destinatário deve ter 11 dígitos (CPF) ou 14 dígitos (CNPJ)."
    ),
    "Recipient CPF cannot be a sequence of repeated digits.": (
        "O CPF do destinatário não pode ser uma sequência de dígitos repetidos."
    ),
    "Recipient CNPJ cannot be a sequence of repeated digits.": (
        "O CNPJ do destinatário não pode ser uma sequência de dígitos repetidos."
    ),
    "Recipient CPF check digit does not match.": (
        "O dígito verificador do CPF do destinatário não confere."
    ),
    "Recipient CNPJ check digit does not match.": (
        "O dígito verificador do CNPJ do destinatário não confere."
    ),
    "Access key must contain 44 digits.": "A chave de acesso deve ter 44 dígitos.",
    "Access key check digit does not match.": (
        "O dígito verificador da chave de acesso não confere."
    ),
    (
        "Access key must contain 44 digits to compare state, year/month, "
        "issuer CNPJ, model, series, and number."
    ): (
        "A chave de acesso deve ter 44 dígitos para comparar UF, ano/mês, "
        "CNPJ do emitente, modelo, série e número."
    ),
    "Access key state does not match state.": (
        "A UF da chave de acesso não confere com a UF."
    ),
    "Access key year/month does not match issue date.": (
        "O ano/mês da chave de acesso não confere com a data de emissão."
    ),
    "Access key issuer CNPJ does not match issuer CNPJ.": (
        "O CNPJ do emitente na chave de acesso não confere com o CNPJ do emitente."
    ),
    "Access key model does not match model.": (
        "O modelo da chave de acesso não confere com o modelo."
    ),
    "Access key series does not match series.": (
        "A série da chave de acesso não confere com a série."
    ),
    "Access key number does not match number.": (
        "O número da chave de acesso não confere com o número."
    ),
    "Item 3: quantity × unit price does not equal the item total within R$ 0.01.": (
        "Item 3: quantidade × valor unitário não é igual ao valor do item, "
        "dentro de R$ 0.01."
    ),
    "Sum of item totals does not equal total products.": (
        "A soma dos valores dos itens não é igual ao total dos produtos."
    ),
    (
        "Invoice total does not equal total products - discount + ICMS-ST total "
        "+ freight + insurance + other expenses + IPI total."
    ): (
        "O valor total da nota não é igual ao total dos produtos − desconto "
        "+ total do ICMS-ST + frete + seguro + outras despesas + total do IPI."
    ),
    "NCM must contain 8 digits.": "O NCM deve ter 8 dígitos.",
    "CFOP must contain 4 digits.": "O CFOP deve ter 4 dígitos.",
    "Issue date is in the future.": "A data de emissão está no futuro.",
    (
        "import tax is an unsupported component; the invoice total rule is skipped "
        "and the document goes straight to human review."
    ): (
        "Imposto de importação é um componente não suportado; a regra do valor "
        "total da nota fica de fora e o documento vai direto para revisão humana."
    ),
    (
        "returned IPI is an unsupported component; the invoice total rule is skipped "
        "and the document goes straight to human review."
    ): (
        "IPI devolvido é um componente não suportado; a regra do valor total da "
        "nota fica de fora e o documento vai direto para revisão humana."
    ),
    (
        "ICMS relief is an unsupported component; the invoice total rule is skipped "
        "and the document goes straight to human review."
    ): (
        "Desoneração do ICMS é um componente não suportado; a regra do valor total "
        "da nota fica de fora e o documento vai direto para revisão humana."
    ),
    (
        "FCP-ST is an unsupported component; the invoice total rule is skipped "
        "and the document goes straight to human review."
    ): (
        "FCP-ST é um componente não suportado; a regra do valor total da nota "
        "fica de fora e o documento vai direto para revisão humana."
    ),
    UPLOAD_RULE: "Envie um PDF, JPG ou PNG.",
    HEIC_RULE: (
        "Fotos HEIC não são aceitas. Exporte a foto como JPG e envie esse arquivo."
    ),
    "This PDF could not be read.": "Não foi possível ler este PDF.",
    "This document has 4 pages, which is over the limit of 3.": (
        "Este documento tem 4 páginas, acima do limite de 3."
    ),
    "This file is larger than 10 MB.": "Este arquivo é maior que 10 MB.",
    "This file is larger than 1500 bytes.": "Este arquivo é maior que 1500 bytes.",
    "This address has already uploaded 1 document today. Try again tomorrow.": (
        "Este endereço já enviou 1 documento hoje. Tente novamente amanhã."
    ),
    "This address has already uploaded 2 documents today. Try again tomorrow.": (
        "Este endereço já enviou 2 documentos hoje. Tente novamente amanhã."
    ),
    "The demo has reached its daily upload limit. Try again tomorrow.": (
        "A demonstração atingiu o limite diário de envios. Tente novamente amanhã."
    ),
    "The demo has reached its daily spend limit. Try again tomorrow.": (
        "A demonstração atingiu o limite diário de custo. Tente novamente amanhã."
    ),
    "The document could not be read.": "Não foi possível ler o documento.",
    "This document does not look like a DANFE.": "Este documento não parece um DANFE.",
    "No document is waiting for review on this thread.": (
        "Nenhum documento aguarda revisão neste envio."
    ),
    "An edit is not a valid value for that field.": (
        "O valor informado não é válido para esse campo."
    ),
    "Only flagged fields can be edited: header.number.": (
        "Só campos sinalizados podem ser editados: header.number."
    ),
    "Feedback could not be saved.": "Não foi possível salvar a sua resposta.",
}


def test_every_validator_and_refusal_message_has_the_approved_portuguese() -> None:
    produced = _validator_messages() | _refusal_messages()

    assert produced == set(_MESSAGES)
    for english, portuguese in _MESSAGES.items():
        assert visitor_text(english, "pt") == portuguese
        assert visitor_text(english, "en") == english


def test_parameterized_messages_keep_the_values_they_carry() -> None:
    assert (
        visitor_text(
            "Item 7: quantity × unit price does not equal the item total within R$ 0.01.",
            "pt",
        )
        == "Item 7: quantidade × valor unitário não é igual ao valor do item, dentro de R$ 0.01."
    )
    assert (
        visitor_text("This document has 9 pages, which is over the limit of 2.", "pt")
        == "Este documento tem 9 páginas, acima do limite de 2."
    )
    assert visitor_text("This file is larger than 3 MB.", "pt") == (
        "Este arquivo é maior que 3 MB."
    )
    assert visitor_text("This file is larger than 2048 bytes.", "pt") == (
        "Este arquivo é maior que 2048 bytes."
    )
    assert visitor_text(
        "This address has already uploaded 5 documents today. Try again tomorrow.",
        "pt",
    ) == ("Este endereço já enviou 5 documentos hoje. Tente novamente amanhã.")
    assert visitor_text(
        "Only flagged fields can be edited: header.invoice_total, line_items[0].quantity.",
        "pt",
    ) == (
        "Só campos sinalizados podem ser editados: "
        "header.invoice_total, line_items[0].quantity."
    )


def test_an_unknown_portuguese_message_is_not_left_in_english() -> None:
    with pytest.raises(KeyError):
        visitor_text("This rule was just invented.", "pt")


def test_intro_uses_the_same_english_as_the_upload_refusals() -> None:
    assert UPLOAD_RULE == UNSUPPORTED_TYPE
    assert HEIC_RULE == HEIC_NOT_ACCEPTED
    portuguese = intro_body("pt")
    english = intro_body("en")
    assert "Envie um PDF, JPG ou PNG." in portuguese
    assert "No máximo 3 páginas e 10 MB." in portuguese
    assert "Fotos HEIC não são aceitas." in portuguese
    assert "não chama o modelo" in portuguese
    assert "Upload a PDF, JPG, or PNG." in english
    assert "At most 3 pages and 10 MB." in english
    assert "does not call a model" in english


def test_field_labels_are_readable_in_both_languages() -> None:
    extraction = {
        "line_items": [{"printed_order": 3}, {"printed_order": 1}],
    }
    labels = {
        "header.access_key": ("Access key", "Chave de acesso"),
        "header.state": ("State", "UF"),
        "header.invoice_total": ("Invoice total", "Valor total da nota"),
        "header.unsupported_components.icms_relief": (
            "ICMS relief",
            "Desoneração do ICMS",
        ),
        "line_items[0].unit_price": ("Item 3 · Unit price", "Item 3 · Valor unitário"),
        "line_items[1].description": ("Item 1 · Description", "Item 1 · Descrição"),
    }
    for path, (english, portuguese) in labels.items():
        assert field_label(path, "en", extraction) == english
        assert field_label(path, "pt", extraction) == portuguese


def test_gallery_captions_translate_the_rendition_and_keep_the_id() -> None:
    assert gallery_caption("synthetic-001 · digital PDF", "pt") == (
        "synthetic-001 · PDF digital"
    )
    assert gallery_caption("synthetic-001 · simulated photo", "en") == (
        "synthetic-001 · simulated photo"
    )
    assert gallery_caption("synthetic-009 · phone photo · planted error", "pt") == (
        "synthetic-009 · Foto de celular · Erro plantado"
    )
    loaded = load_gallery(_SYNTHETIC, today=_TODAY)
    portuguese = {gallery_caption(item.caption, "pt") for item in loaded}
    assert "synthetic-001 · PDF digital" in portuguese
    assert "synthetic-001 · Foto simulada" in portuguese
    assert any(caption.endswith("· Erro plantado") for caption in portuguese)
    assert all("planted error" not in caption for caption in portuguese)
    assert all("digital PDF" not in caption for caption in portuguese)


def test_invoice_data_appears_as_read_in_both_languages() -> None:
    body = next(
        item.payload
        for item in load_gallery(_SYNTHETIC, today=_TODAY)
        if item.caption == "synthetic-001 · digital PDF"
    )
    portuguese = field_rows(body, language="pt")
    english = field_rows(body, language="en")

    assert [row[2] for row in portuguese] == [row[2] for row in english]
    values = {row[2] for row in portuguese}
    assert "83606369122950" in values
    assert "2024-09-07" in values
    assert any(value.startswith("Parafuso sextavado") for value in values)
    assert excel_filename(body) == "nfe-1-1.xlsx"
    pt_book = load_workbook(io.BytesIO(workbook_bytes(body, language="pt")))
    en_book = load_workbook(io.BytesIO(workbook_bytes(body, language="en")))
    assert _column(pt_book["Cabeçalho"], 2) == _column(en_book["Header"], 2)
    assert _column(pt_book["Itens"], 2) == _column(en_book["Line items"], 2)


def test_the_page_opens_in_portuguese_and_english_redraws_without_a_model(
    tmp_path,
) -> None:
    reader = ScriptedReader(_extraction(issuer_cnpj="00000000000181"))
    actions = DemoActions(
        _service(reader, RecordingMonitor()),
        load_gallery(_SYNTHETIC, today=_TODAY),
    )
    opening = language_chrome(actions, "Português")
    assert opening.language == "pt"
    assert "Envie um PDF, JPG ou PNG." in opening.intro
    assert "synthetic-001 · PDF digital" in opening.captions
    assert reader.requests == []

    english = language_chrome(actions, "English")
    assert english.language == "en"
    assert "Upload a PDF, JPG, or PNG." in english.intro
    assert "synthetic-001 · digital PDF" in english.captions
    assert reader.requests == []

    screen = present_upload(actions, _pdf_path(tmp_path), client_ip=_IP)
    calls = len(reader.requests)
    portuguese = screen_copy(screen, "pt")
    selected = screen_copy(screen, "en")
    again = screen_copy(screen, "pt")

    assert len(reader.requests) == calls
    assert portuguese.outcome_line == "**Resultado:** Precisa de revisão"
    assert selected.outcome_line == "**Outcome:** Needs review"
    assert portuguese.feedback_question == "Esta leitura está correta?"
    assert selected.feedback_question == "Was this reading correct?"
    assert portuguese.decimal_hint == "Use ponto como separador decimal."
    assert selected.decimal_hint is None
    assert _labels(portuguese) == _labels(again)
    issuer = next(
        field
        for field in portuguese.review_fields
        if field.path == "header.issuer_cnpj"
    )
    issuer_en = next(
        field for field in selected.review_fields if field.path == "header.issuer_cnpj"
    )
    assert issuer.label == "CNPJ do emitente"
    assert issuer_en.label == "Issuer CNPJ"
    assert issuer.value == issuer_en.value == "00000000000181"
    assert [row[2] for row in portuguese.rows] == [row[2] for row in selected.rows]


def test_a_language_change_rebuilds_review_boxes(
    tmp_path,
) -> None:
    actions = DemoActions(
        _service(
            ScriptedReader(_extraction(issuer_cnpj="00000000000181")),
            RecordingMonitor(),
        ),
        (),
    )
    screen = present_upload(actions, _pdf_path(tmp_path), client_ip=_IP)
    portuguese = _review_keys(actions, screen, "pt")
    english = _review_keys(actions, screen, "en")
    assert portuguese and english
    assert {path for _, path in portuguese} == {path for _, path in english}
    assert {language for language, _ in portuguese} == {"pt"}
    assert {language for language, _ in english} == {"en"}
    assert ("pt", "header.issuer_cnpj") in portuguese


def _review_keys(
    actions: DemoActions, screen: dict[str, object], language: str
) -> set[tuple[str, str]]:
    with gr.Blocks() as demo:
        state = gr.State(screen)
        render_result(actions, screen, state, language)
    return {
        block.key
        for block in demo.blocks.values()
        if isinstance(block, gr.Textbox) and isinstance(block.key, tuple)
    }


def test_language_choice_is_stored_in_the_browser_script() -> None:
    actions = DemoActions(
        _service(ScriptedReader(_extraction()), RecordingMonitor()), ()
    )
    demo = build_blocks(actions)
    radio = next(block for block in demo.blocks.values() if isinstance(block, gr.Radio))
    assert radio.value == "Português"
    config = json.dumps(demo.config, default=str)
    assert "nfe-demo-language" in config
    assert "localStorage.getItem" in config
    assert "localStorage.setItem" in config
    assert "navigator.language" not in config
    assert "Portugu" in config
    store = next(
        dep for dep in demo.config["dependencies"] if dep["api_name"] == "_apply"
    )
    restore = next(
        dep for dep in demo.config["dependencies"] if dep["api_name"] == "_restore"
    )
    render = next(dep for dep in demo.config["dependencies"] if dep["render_id"] == 0)
    assert store["queue"] is True
    assert restore["queue"] is True
    language_state = store["outputs"][0]
    assert (language_state, "change") in render["targets"]


def _labels(view: ScreenCopy) -> tuple[str, ...]:
    return tuple(field.label for field in view.review_fields)


def _column(sheet: object, index: int) -> list[object]:
    rows = list(sheet.iter_rows(min_row=2, values_only=True))  # type: ignore[attr-defined]
    return [row[index - 1] for row in rows]


def _pdf_path(folder, name: str = "invoice.pdf") -> str:
    path = folder / name
    path.write_bytes(_pdf(1))
    return str(path)


def _validator_messages() -> set[str]:
    found: set[str] = set()
    cases = (
        _extraction(issuer_cnpj="52998224725"),
        _extraction(issuer_cnpj="00000000000000"),
        _extraction(issuer_cnpj="00000000000181"),
        _extraction(recipient_cnpj_cpf="12345"),
        _extraction(recipient_cnpj_cpf="00000000000"),
        _extraction(recipient_cnpj_cpf="00000000000000"),
        _extraction(recipient_cnpj_cpf="52998224726"),
        _extraction(recipient_cnpj_cpf="11222333000182"),
        _extraction(access_key="123"),
        _extraction(access_key=ACCESS_KEY[:-1] + "1"),
        _extraction(state="RJ"),
        _extraction(model="65"),
        _extraction(series="2"),
        _extraction(number="124"),
        _extraction(issue_date=date(2024, 11, 3)),
        _extraction(
            line_items=(
                _line_item(
                    printed_order=3,
                    quantity=Decimal(9),
                    unit_price=Decimal("8.00"),
                    ncm="123",
                    cfop="99",
                ),
            )
        ),
        _extraction(total_products=2_000),
        _extraction(invoice_total=1),
        _extraction(
            unsupported_components={
                "import_tax": True,
                "returned_ipi": True,
                "icms_relief": True,
                "fcp_st": True,
            }
        ),
    )
    for extraction in cases:
        report = validate(extraction, today=_TODAY)
        found.update(
            failure.message for rule in report.rules for failure in rule.failures
        )
        found.update(report.review_messages)
    future = validate(_extraction(), today=date(2024, 10, 2))
    found.update(failure.message for rule in future.rules for failure in rule.failures)
    return found


def _refusal_messages() -> set[str]:
    found: set[str] = set()
    found.add(_refused(_too_large_bytes))
    found.add(_refused(_too_large_megabytes))
    found.add(_refused(_heic))
    found.add(_refused(_wrong_type))
    found.add(_refused(_long_pdf))
    found.add(_refused(_broken_pdf))
    found.update(_service_refusals())
    return found


def _too_large_bytes() -> None:
    prepare_upload(
        b"x" * 1501,
        filename="a.pdf",
        content_type="application/pdf",
        limits=CostLimits(max_bytes=1500),
    )


def _too_large_megabytes() -> None:
    prepare_upload(
        b"x" * (10 * 1024 * 1024 + 1),
        filename="a.pdf",
        content_type="application/pdf",
        limits=CostLimits(),
    )


def _heic() -> None:
    prepare_upload(
        _HEIC,
        filename="photo.heic",
        content_type="image/heic",
        limits=CostLimits(),
    )


def _wrong_type() -> None:
    prepare_upload(
        b"notes",
        filename="notes.txt",
        content_type="text/plain",
        limits=CostLimits(),
    )


def _long_pdf() -> None:
    prepare_upload(
        _pdf(4),
        filename="long.pdf",
        content_type="application/pdf",
        limits=CostLimits(),
    )


def _broken_pdf() -> None:
    prepare_upload(
        b"%PDF-1.4 not a pdf",
        filename="broken.pdf",
        content_type="application/pdf",
        limits=CostLimits(),
    )


def _refused(call) -> str:
    with pytest.raises(RequestRejected) as caught:
        call()
    return caught.value.message


def _service_refusals() -> set[str]:
    found: set[str] = set()
    unreadable = DemoActions(
        _service(_UnreadableReader(), RecordingMonitor()),
        (),
    )
    refused = unreadable.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip=_IP,
    )
    assert hasattr(refused, "message")
    found.add(refused.message)  # type: ignore[union-attr]

    rejected = DemoActions(
        _service(ScriptedReader(_extraction(access_key="")), RecordingMonitor()),
        (),
    )
    not_danfe = rejected.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip=_IP,
    )
    found.add(not_danfe.message)  # type: ignore[union-attr]

    review = DemoActions(
        _service(
            ScriptedReader(_extraction(issuer_cnpj="00000000000181")),
            RecordingMonitor(),
        ),
        (),
    )
    uploaded = review.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip=_IP,
    )
    assert isinstance(uploaded, dict)
    thread_id = uploaded["thread_id"]
    assert isinstance(thread_id, str)
    invalid = review.submit_review(thread_id, {"header.issuer_cnpj": None})
    only_flagged = review.submit_review(thread_id, {"header.number": "999"})
    missing = review.submit_review("missing-thread", {})
    found.add(invalid.message)  # type: ignore[union-attr]
    found.add(only_flagged.message)  # type: ignore[union-attr]
    found.add(missing.message)  # type: ignore[union-attr]

    singular = DemoActions(
        _service(
            ScriptedReader(_extraction()),
            RecordingMonitor(),
            limits=CostLimits(uploads_per_ip_per_day=1),
        ),
        (),
    )
    assert isinstance(
        singular.upload(
            _pdf(1), filename="a.pdf", content_type="application/pdf", client_ip=_IP
        ),
        dict,
    )
    blocked = singular.upload(
        _pdf(1),
        filename="b.pdf",
        content_type="application/pdf",
        client_ip=_IP,
    )
    found.add(blocked.message)  # type: ignore[union-attr]

    plural = DemoActions(
        _service(
            ScriptedReader(_extraction()),
            RecordingMonitor(),
            limits=CostLimits(uploads_per_ip_per_day=2),
        ),
        (),
    )
    for name, address in (("a.pdf", "198.51.100.31"), ("b.pdf", "198.51.100.31")):
        assert isinstance(
            plural.upload(
                _pdf(1),
                filename=name,
                content_type="application/pdf",
                client_ip=address,
            ),
            dict,
        )
    plural_block = plural.upload(
        _pdf(1),
        filename="c.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.31",
    )
    found.add(plural_block.message)  # type: ignore[union-attr]

    shared = DemoActions(
        _service(
            ScriptedReader(_extraction()),
            RecordingMonitor(),
            limits=CostLimits(global_uploads_per_day=1),
        ),
        (),
    )
    assert isinstance(
        shared.upload(
            _pdf(1),
            filename="a.pdf",
            content_type="application/pdf",
            client_ip="198.51.100.41",
        ),
        dict,
    )
    global_cap = shared.upload(
        _pdf(1),
        filename="b.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.42",
    )
    found.add(global_cap.message)  # type: ignore[union-attr]

    spender = DemoActions(
        _service(
            ScriptedReader(_extraction(), output_tokens=_OVER_CAP_OUTPUT_TOKENS),
            RecordingMonitor(),
        ),
        (),
    )
    assert isinstance(
        spender.upload(
            _pdf(1),
            filename="a.pdf",
            content_type="application/pdf",
            client_ip="198.51.100.51",
        ),
        dict,
    )
    spend_cap = spender.upload(
        _pdf(1),
        filename="b.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.52",
    )
    found.add(spend_cap.message)  # type: ignore[union-attr]

    feedback = present_visitor_feedback(
        review,
        {"feedback_ref": "", "body": uploaded},
        score=1,
        comment="",
    )
    notice = feedback.get("feedback_notice")
    assert isinstance(notice, str)
    found.add(notice)
    return found


class _UnreadableReader:
    def read(self, request: SectionRequest) -> SectionReading:
        return SectionReading(
            section=request.section,
            failure=ExtractionFailure.PARSE_FAILURE,
        )
