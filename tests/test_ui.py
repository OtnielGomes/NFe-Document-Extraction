"""UI seam: gallery, upload, and review over the same extraction service.

The model reader and the demo monitor are the boundaries. Gallery results
are cached readings of the printed synthetic invoices.
"""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote

import gradio as gr
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_api import RecordingMonitor, ScriptedReader, _pdf
from test_identifier_rules import ISSUER_CNPJ, _extraction, _line_item

from nfe_document_extraction.agents.config import CHEAP_VISION, load_approach_config
from nfe_document_extraction.agents.sections import (
    SectionReader,
    SectionReading,
    SectionRequest,
)
from nfe_document_extraction.app.api import create_app
from nfe_document_extraction.app.limits import CostLimits
from nfe_document_extraction.app.service import ExtractionService
from nfe_document_extraction.app.ui import (
    DemoActions,
    Refusal,
    document_panel,
    editable_field_paths,
    field_rows,
    gradio_client_ip,
    load_gallery,
    present_gallery_selection,
    present_review,
    present_upload,
    review_inputs,
)

_TODAY = date(2026, 10, 6)
_SYNTHETIC = Path("data/synthetic")
_IP = "198.51.100.10"
_HEIC = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00"
# 1_250_000 output tokens at the cheap-vision price is US$ 0.625.
_OVER_CAP_OUTPUT_TOKENS = 1_250_000


class _ExplodingReader:
    """Fails the test if the gallery reaches the model."""

    def read(self, request: SectionRequest) -> SectionReading:
        raise AssertionError(f"gallery called the model for {request.section}")


def test_a_gallery_item_shows_the_cached_reading_without_a_model_call_or_a_trace() -> (
    None
):
    reader = _ExplodingReader()
    monitor = RecordingMonitor()
    actions = DemoActions(
        _service(reader, monitor),
        load_gallery(_SYNTHETIC, today=_TODAY),
    )

    cached = _open(actions, "synthetic-001 · digital PDF")
    again = _open(actions, "synthetic-001 · digital PDF")
    photo = _open(actions, "synthetic-001 · simulated photo")
    planted = _open(actions, "synthetic-007 · digital PDF")

    assert cached == again
    assert cached["outcome"] == "accepted"
    assert cached["thread_id"] is None
    assert cached["review_messages"] == []
    assert _field(cached, "header.issuer_cnpj") == {
        "field": "header.issuer_cnpj",
        "status": "verified",
        "messages": [],
    }
    assert cached["extraction"]["header"]["issuer_cnpj"] == "83606369122950"
    assert _field(cached, "line_items[0].description")["status"] == "unverifiable"
    assert photo["outcome"] == "accepted"
    assert photo["extraction"] == cached["extraction"]
    assert photo["thread_id"] is None
    invoice_total = _field(planted, "header.invoice_total")
    assert planted["outcome"] == "human_review"
    assert planted["thread_id"] is None
    assert invoice_total["status"] == "flagged"
    assert any(
        "Invoice total does not equal" in message
        for message in invoice_total["messages"]
    )
    assert editable_field_paths(planted) == ()
    assert monitor.extractions == []
    assert monitor.corrections == []


def test_an_upload_marks_each_field_verified_unverifiable_or_flagged() -> None:
    reader = ScriptedReader(_extraction(issuer_cnpj="00000000000181"))
    actions = DemoActions(_service(reader, RecordingMonitor()), ())

    uploaded = actions.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.10",
    )

    assert isinstance(uploaded, dict)
    rows = {row[0]: row for row in field_rows(uploaded)}
    assert rows["header.issuer_cnpj"][1] == "Flagged"
    assert rows["header.issuer_cnpj"][2] == "00000000000181"
    assert "Issuer CNPJ check digit does not match." in rows["header.issuer_cnpj"][3]
    assert rows["line_items[0].description"] == [
        "line_items[0].description",
        "Unverifiable",
        "Plain bearing",
        "",
    ]
    verified = next(
        field["field"]
        for field in uploaded["fields"]
        if isinstance(field, dict) and field["status"] == "verified"
    )
    assert isinstance(verified, str)
    assert rows[verified][1] == "Verified"
    assert rows[verified][3] == ""


def test_human_review_edits_only_flagged_fields_and_revalidates() -> None:
    reader = ScriptedReader(_extraction(issuer_cnpj="00000000000181"))
    actions = DemoActions(_service(reader, RecordingMonitor()), ())

    uploaded = actions.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.10",
    )

    assert isinstance(uploaded, dict)
    assert uploaded["outcome"] == "human_review"
    thread_id = uploaded["thread_id"]
    assert isinstance(thread_id, str) and thread_id
    editable = editable_field_paths(uploaded)
    flagged = {
        field["field"]
        for field in uploaded["fields"]
        if isinstance(field, dict) and field["status"] == "flagged"
    }
    assert isinstance(flagged, set)
    assert set(editable) == flagged
    assert "header.issuer_cnpj" in editable
    assert "line_items[0].description" not in editable
    assert "line_items[0].product_code" not in editable
    calls = len(reader.requests)

    reviewed = actions.submit_review(thread_id, {"header.issuer_cnpj": ISSUER_CNPJ})

    assert isinstance(reviewed, dict)
    assert reviewed["outcome"] == "accepted"
    assert reviewed["thread_id"] is None
    assert reviewed["extraction"]["header"]["issuer_cnpj"] == ISSUER_CNPJ
    assert _field(reviewed, "header.issuer_cnpj")["status"] == "verified"
    assert _field(reviewed, "header.issuer_cnpj")["messages"] == []
    assert editable_field_paths(reviewed) == ()
    assert len(reader.requests) == calls
    shown = review_inputs(uploaded)
    assert [item.path for item in shown] == list(editable)
    issuer = next(item for item in shown if item.path == "header.issuer_cnpj")
    assert issuer.value == "00000000000181"


def test_review_keeps_money_when_the_shown_value_is_submitted() -> None:
    reader = ScriptedReader(
        _extraction(issuer_cnpj="00000000000181", invoice_total=999)
    )
    actions = DemoActions(_service(reader, RecordingMonitor()), ())
    uploaded = actions.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.10",
    )

    assert isinstance(uploaded, dict)
    thread_id = uploaded["thread_id"]
    assert isinstance(thread_id, str)
    edits = {item.path: item.value for item in review_inputs(uploaded)}
    assert edits["header.invoice_total"] == "9.99"
    assert edits["header.freight"] == "0.00"
    edits["header.issuer_cnpj"] = ISSUER_CNPJ

    reviewed = actions.submit_review(thread_id, edits)

    assert isinstance(reviewed, dict)
    header = reviewed["extraction"]["header"]
    assert header["invoice_total"] == 999
    assert header["freight"] == 0
    assert header["total_products"] == 1_000
    assert header["icms_st_total"] == 180
    assert reviewed["extraction"]["line_items"][0]["item_total"] == 1_000


def test_an_upload_with_no_access_key_says_it_is_not_a_danfe(tmp_path: Path) -> None:
    reader = ScriptedReader(_extraction(access_key=""))
    actions = DemoActions(_service(reader, RecordingMonitor()), ())

    shown = present_upload(actions, _write_pdf(tmp_path), client_ip=_IP)

    assert shown["notice"] == "This document does not look like a DANFE."
    assert shown["body"] is None


def test_the_screen_shows_upload_rules_refusals_and_review_errors(
    tmp_path: Path,
) -> None:
    reader = ScriptedReader(_extraction(issuer_cnpj="00000000000181"))
    actions = DemoActions(_service(reader, RecordingMonitor()), ())
    missing = present_upload(actions, None, client_ip=_IP)
    assert missing == {"notice": "Choose a PDF, JPG, or PNG.", "body": None}

    uploaded = present_upload(actions, _write_pdf(tmp_path), client_ip=_IP)
    assert uploaded["notice"] is None
    body = uploaded["body"]
    assert isinstance(body, dict)
    assert body["outcome"] == "human_review"
    kept = present_review(actions, uploaded, {"header.number": "999"})
    assert kept["notice"] == "Only flagged fields can be edited: header.number."
    assert kept["body"] is body
    done = present_review(actions, kept, {"header.issuer_cnpj": ISSUER_CNPJ})
    assert done["notice"] is None
    assert isinstance(done["body"], dict)
    assert done["body"]["outcome"] == "accepted"


def test_the_demo_page_and_the_api_share_one_upload_limit(tmp_path: Path) -> None:
    reader = ScriptedReader(_extraction())
    gallery = load_gallery(_SYNTHETIC, today=_TODAY)
    app = create_app(
        reader,
        monitor=RecordingMonitor(),
        limits=CostLimits(uploads_per_ip_per_day=1),
        clock=lambda: _TODAY,
        gallery=gallery,
    )
    actions = app.state.demo_actions
    assert isinstance(actions, DemoActions)
    opened = present_gallery_selection(actions, 0)
    assert opened["notice"] is None
    assert isinstance(opened["body"], dict)
    assert opened["body"]["outcome"] == "accepted"
    assert present_gallery_selection(actions, 99)["notice"] == (
        "Choose an invoice from the gallery."
    )
    assert reader.requests == []

    with TestClient(app) as client:
        page = client.get("/")
        blocked_setup = present_upload(actions, _write_pdf(tmp_path), client_ip=_IP)
        assert blocked_setup["notice"] is None
        blocked = client.post(
            "/extract",
            files={"document": ("invoice.pdf", _pdf(1), "application/pdf")},
            headers={"X-Forwarded-For": _IP},
        )

    assert page.status_code == 200
    text = page.text + json.dumps(app.state.demo_blocks.config, default=str)
    assert "NFe Document Extraction" in text
    assert "At most 3 pages and 10 MB." in text
    assert (
        "HEIC photos are not accepted. Export the photo as JPG and upload that file."
        in text
    )
    assert "does not call a model" in text
    galleries = [
        component
        for component in app.state.demo_blocks.config["components"]
        if component.get("type") == "gallery"
    ]
    assert galleries[0]["props"]["sources"] == []
    assert galleries[0]["props"]["interactive"] is False
    assert blocked.status_code == 429
    assert blocked.json()["detail"] == (
        "This address has already uploaded 1 document today. Try again tomorrow."
    )


class _GradioRequest:
    def __init__(self, forwarded: str | None, peer: str | None) -> None:
        self.headers = {} if forwarded is None else {"x-forwarded-for": forwarded}
        self.client = None if peer is None else _Peer(peer)


class _Peer:
    def __init__(self, host: str) -> None:
        self.host = host


def test_a_gradio_request_uses_the_same_visitor_address_as_the_api() -> None:
    assert (
        gradio_client_ip(_GradioRequest("203.0.113.8, 10.0.0.4", "10.0.0.8"))
        == "203.0.113.8"
    )
    assert gradio_client_ip(_GradioRequest(None, "10.0.0.8")) == "10.0.0.8"
    assert gradio_client_ip(_GradioRequest("  ", None)) == "unknown"


def test_upload_rules_and_daily_limits_are_shown_in_plain_language() -> None:
    reader = ScriptedReader(_extraction())
    actions = DemoActions(
        _service(
            reader, RecordingMonitor(), limits=CostLimits(uploads_per_ip_per_day=1)
        ),
        (),
    )

    oversized = _upload(actions, b"x" * (10 * 1024 * 1024 + 1), "big.pdf")
    heic = _upload(actions, _HEIC, "photo.heic")
    long_pdf = _upload(actions, _pdf(4), "long.pdf")
    wrong_type = _upload(actions, b"not an invoice", "notes.txt")

    assert isinstance(oversized, Refusal)
    assert oversized.message == "This file is larger than 10 MB."
    assert isinstance(heic, Refusal)
    assert heic.message == (
        "HEIC photos are not accepted. Export the photo as JPG and upload that file."
    )
    assert isinstance(long_pdf, Refusal)
    assert long_pdf.message == (
        "This document has 4 pages, which is over the limit of 3."
    )
    assert isinstance(wrong_type, Refusal)
    assert wrong_type.message == "Upload a PDF, JPG, or PNG."
    assert reader.requests == []

    assert isinstance(_upload(actions, _pdf(1), "invoice.pdf"), dict)
    blocked = _upload(actions, _pdf(1), "another.pdf")
    assert isinstance(blocked, Refusal)
    assert blocked.message == (
        "This address has already uploaded 1 document today. Try again tomorrow."
    )
    calls = len(reader.requests)

    shared = DemoActions(
        _service(
            reader,
            RecordingMonitor(),
            limits=CostLimits(global_uploads_per_day=1),
        ),
        (),
    )
    assert isinstance(_upload(shared, _pdf(1), "a.pdf", ip="198.51.100.1"), dict)
    global_cap = _upload(shared, _pdf(1), "b.pdf", ip="198.51.100.2")
    assert isinstance(global_cap, Refusal)
    assert global_cap.message == (
        "The demo has reached its daily upload limit. Try again tomorrow."
    )

    spender = ScriptedReader(_extraction(), output_tokens=_OVER_CAP_OUTPUT_TOKENS)
    spend = DemoActions(_service(spender, RecordingMonitor()), ())
    assert isinstance(_upload(spend, _pdf(1), "invoice.pdf"), dict)
    spend_cap = _upload(spend, _pdf(1), "invoice.pdf", ip="198.51.100.11")
    assert isinstance(spend_cap, Refusal)
    assert spend_cap.message == (
        "The demo has reached its daily spend limit. Try again tomorrow."
    )
    assert len(reader.requests) == calls + 2


def _write_pdf(folder: Path) -> str:
    path = folder / "invoice.pdf"
    path.write_bytes(_pdf(1))
    return str(path)


def _upload(
    actions: DemoActions,
    data: bytes,
    filename: str,
    *,
    ip: str = _IP,
) -> dict[str, object] | Refusal:
    return actions.upload(
        data,
        filename=filename,
        content_type="application/pdf",
        client_ip=ip,
    )


def _service(
    reader: SectionReader,
    monitor: RecordingMonitor,
    limits: CostLimits | None = None,
) -> ExtractionService:
    return ExtractionService(
        reader,
        monitor=monitor,
        limits=CostLimits() if limits is None else limits,
        clock=lambda: _TODAY,
        price=load_approach_config(CHEAP_VISION),
    )


def test_the_page_offers_excel_for_gallery_uploads_and_reviewed_results() -> None:
    gallery = load_gallery(_SYNTHETIC, today=_TODAY)
    actions = DemoActions(
        _service(
            ScriptedReader(_extraction(line_items=(_line_item(quantity=Decimal(9)),))),
            RecordingMonitor(),
        ),
        gallery,
    )
    uploaded = actions.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip=_IP,
    )
    assert isinstance(uploaded, dict)
    thread_id = uploaded["thread_id"]
    assert isinstance(thread_id, str)
    edits = {item.path: item.value for item in review_inputs(uploaded)}
    edits["line_items[0].quantity"] = "1.2500"
    reviewed = actions.submit_review(thread_id, edits)
    assert isinstance(reviewed, dict)

    offered = (
        (actions.open_gallery_item(0), "nfe-1-1.xlsx"),
        (uploaded, "nfe-123-1.xlsx"),
        (reviewed, "nfe-123-1.xlsx"),
    )
    for body, filename in offered:
        button = _excel_button(actions, body)
        assert button.label == "Download Excel"
        assert isinstance(button.value, dict)
        assert button.value["orig_name"] == filename
        book = load_workbook(_download_file(button.value))
        assert book.sheetnames == ["Summary", "Header", "Line items"]


def _excel_button(actions: DemoActions, body: dict[str, object]) -> gr.DownloadButton:
    with gr.Blocks() as demo:
        screen = gr.State({"notice": None, "body": body})
        document_panel(actions, body, screen)
    dataframes = [
        block for block in demo.blocks.values() if isinstance(block, gr.Dataframe)
    ]
    buttons = [
        block for block in demo.blocks.values() if isinstance(block, gr.DownloadButton)
    ]
    assert len(dataframes) == 1
    assert len(buttons) == 1
    assert dataframes[0]._id < buttons[0]._id
    return buttons[0]


def _download_file(value: dict[str, object]) -> io.BytesIO | Path:
    raw = value.get("path")
    assert isinstance(raw, str)
    path = Path(unquote(raw))
    if path.is_file():
        return path
    raise AssertionError(raw)


def test_demo_client_starts_without_an_api_key(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    from nfe_document_extraction.app.api import _openai_client

    client = _openai_client()
    assert client.api_key == "missing"


def _open(actions: DemoActions, caption: str) -> dict[str, object]:
    for index, card in enumerate(actions.gallery_cards()):
        if card[1] == caption or card[1].startswith(caption + " ·"):
            body = actions.open_gallery_item(index)
            assert isinstance(body, dict)
            return body
    raise AssertionError(caption)


def _field(body: dict[str, object], name: str) -> dict[str, object]:
    fields = body["fields"]
    assert isinstance(fields, list)
    for field in fields:
        assert isinstance(field, dict)
        if field["field"] == name:
            return field
    raise AssertionError(name)
