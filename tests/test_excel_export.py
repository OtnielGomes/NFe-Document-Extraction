"""Excel export of the document the demo already holds. No Gradio."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from test_api import RecordingMonitor, ScriptedReader, _pdf
from test_identifier_rules import _extraction, _line_item

from nfe_document_extraction.agents.config import CHEAP_VISION, load_approach_config
from nfe_document_extraction.app.excel import excel_filename, workbook_bytes
from nfe_document_extraction.app.gallery import load_gallery
from nfe_document_extraction.app.limits import CostLimits
from nfe_document_extraction.app.service import ExtractionService, document_payload
from nfe_document_extraction.app.ui import DemoActions, review_inputs
from nfe_document_extraction.domain import Extraction
from nfe_document_extraction.validation import validate
from nfe_document_extraction.validation.report import DocumentOutcome

_TODAY = date(2026, 10, 6)
_SYNTHETIC = Path("data/synthetic")
_DESCRIPTION = (
    "Parafuso sextavado zincado com arruela de pressao e porca "
    "para fixacao de estrutura metalica"
)


def test_an_accepted_document_exports_three_sheets_with_typed_cells() -> None:
    body = _gallery_body("synthetic-001 · digital PDF")

    assert excel_filename(body) == "nfe-1-1.xlsx"
    book = load_workbook(io.BytesIO(workbook_bytes(body, language="en")))

    assert book.sheetnames == ["Summary", "Header", "Line items"]
    summary = book["Summary"]
    assert summary["A1"].value == "Outcome"
    assert summary["B1"].value == "Accepted"
    assert summary["A2"].value == "Review messages"
    assert summary["B2"].value is None

    header = book["Header"]
    assert [cell.value for cell in header[1]] == [
        "Field",
        "Value",
        "Status",
        "Messages",
    ]
    issuer = _header_row(header, "Issuer CNPJ")
    assert issuer[1].value == "83606369122950"
    assert issuer[1].data_type == "s"
    assert issuer[1].number_format == "@"
    assert issuer[2].value == "Verified"
    assert issuer[3].value is None
    recipient = _header_row(header, "Recipient CNPJ/CPF")
    assert recipient[1].value == "00149687169118"
    assert recipient[1].data_type == "s"
    access_key = _header_row(header, "Access key")
    assert access_key[1].value == "35240983606369122950550010000000011645576624"
    assert access_key[1].data_type == "s"
    issued = _header_row(header, "Issue date")
    assert issued[1].is_date
    issued_on = issued[1].value
    assert isinstance(issued_on, datetime)
    assert (issued_on.year, issued_on.month, issued_on.day) == (2024, 9, 7)
    products = _header_row(header, "Total products")
    assert products[1].number_format == "0.00"
    assert f"{products[1].value:.2f}" == "44579.28"
    invoice = _header_row(header, "Invoice total")
    assert invoice[1].number_format == "0.00"
    assert f"{invoice[1].value:.2f}" == "52157.76"
    freight = _header_row(header, "Freight")
    assert freight[1].number_format == "0.00"
    assert f"{freight[1].value:.2f}" == "0.00"
    assert _header_row(header, "Import tax")[2].value == ("Unverifiable")

    items = book["Line items"]
    assert [cell.value for cell in items[1]] == [
        "Product code",
        "Description",
        "NCM",
        "CFOP",
        "Quantity",
        "Unit price",
        "Item total",
    ]
    extraction = body["extraction"]
    assert isinstance(extraction, dict)
    line_items = extraction["line_items"]
    assert isinstance(line_items, list)
    assert items.max_row == len(line_items) + 1
    first = _item_row(items, "CAM-01")
    assert first[1].value == _DESCRIPTION
    assert first[2].value == "61091000"
    assert first[2].data_type == "s"
    assert first[2].number_format == "@"
    assert first[3].value == "5405"
    assert first[3].data_type == "s"
    assert first[4].number_format == "0.0000"
    assert f"{first[4].value:.4f}" == "11.1933"
    assert first[5].number_format == "0.0000"
    assert f"{first[5].value:.4f}" == "197.7105"
    assert first[6].number_format == "0.00"
    assert f"{first[6].value:.2f}" == "2213.03"
    coffee = _item_row(items, "CAF-02")
    assert coffee[2].value == "09012100"
    assert coffee[2].data_type == "s"


def test_a_document_in_needs_review_highlights_flagged_line_cells() -> None:
    body = _payload(
        _extraction(
            line_items=(
                _line_item(
                    ncm="09012100",
                    cfop="0102",
                    quantity=Decimal(9),
                    unit_price=Decimal("8.00"),
                ),
            ),
            series="001",
            number="000000123",
            unsupported_components={
                "import_tax": True,
                "returned_ipi": False,
                "icms_relief": False,
                "fcp_st": False,
            },
        )
    )

    assert excel_filename(body) == "nfe-000000123-001.xlsx"
    book = load_workbook(io.BytesIO(workbook_bytes(body, language="en")))
    summary = book["Summary"]
    assert summary["B1"].value == "Needs review"
    assert summary["B2"].value == (
        "import tax is an unsupported component; the invoice total rule is skipped "
        "and the document goes straight to human review."
    )
    header = book["Header"]
    series = _header_row(header, "Series")
    assert series[1].value == "001"
    assert series[1].data_type == "s"
    number = _header_row(header, "Number")
    assert number[1].value == "000000123"
    assert number[1].data_type == "s"
    issuer = _header_row(header, "Issuer CNPJ")
    assert issuer[1].value == "00000000000191"
    assert issuer[1].data_type == "s"
    access_key = _header_row(header, "Access key")
    assert access_key[1].value == "35241000000000000191550010000001231000000010"
    assert access_key[1].data_type == "s"

    row = _item_row(book["Line items"], "SKU-1")
    assert row[2].value == "09012100"
    assert row[2].data_type == "s"
    assert row[3].value == "0102"
    assert row[3].data_type == "s"
    assert _highlighted(row[0]) is False
    assert _highlighted(row[1]) is False
    assert _highlighted(row[2]) is False
    assert _highlighted(row[3]) is False
    assert _highlighted(row[4]) is True
    assert _highlighted(row[5]) is True
    assert _highlighted(row[6]) is True
    assert row[4].number_format == "0"
    assert f"{row[4].value:.0f}" == "9"
    assert row[5].number_format == "0.00"
    assert f"{row[5].value:.2f}" == "8.00"
    assert row[6].number_format == "0.00"
    assert f"{row[6].value:.2f}" == "10.00"


def test_a_reviewed_document_exports_the_corrected_line_without_highlights() -> None:
    actions = DemoActions(
        ExtractionService(
            ScriptedReader(_extraction(line_items=(_line_item(quantity=Decimal(9)),))),
            monitor=RecordingMonitor(),
            limits=CostLimits(),
            clock=lambda: _TODAY,
            price=load_approach_config(CHEAP_VISION),
        ),
        (),
    )
    uploaded = actions.upload(
        _pdf(1),
        filename="invoice.pdf",
        content_type="application/pdf",
        client_ip="198.51.100.10",
    )
    assert isinstance(uploaded, dict)
    assert uploaded["outcome"] == "human_review"
    thread_id = uploaded["thread_id"]
    assert isinstance(thread_id, str)
    edits = {item.path: item.value for item in review_inputs(uploaded)}
    edits["line_items[0].quantity"] = "1.2500"

    reviewed = actions.submit_review(thread_id, edits)

    assert isinstance(reviewed, dict)
    assert reviewed["outcome"] == "accepted"
    assert reviewed["thread_id"] is None
    assert excel_filename(reviewed) == "nfe-123-1.xlsx"
    book = load_workbook(io.BytesIO(workbook_bytes(reviewed, language="en")))
    assert book["Summary"]["B1"].value == "Accepted"
    assert book["Summary"]["B2"].value is None
    row = _item_row(book["Line items"], "SKU-1")
    assert _highlighted(row[4]) is False
    assert _highlighted(row[5]) is False
    assert _highlighted(row[6]) is False
    assert row[4].number_format == "0.0000"
    assert f"{row[4].value:.4f}" == "1.2500"
    assert row[6].number_format == "0.00"
    assert f"{row[6].value:.2f}" == "10.00"


def test_filename_falls_back_when_number_or_series_cannot_be_read() -> None:
    body = _gallery_body("synthetic-001 · digital PDF")
    assert (
        excel_filename(_without_header_value(body, "number")) == "nfe-extraction.xlsx"
    )
    assert (
        excel_filename(_without_header_value(body, "series")) == "nfe-extraction.xlsx"
    )
    assert excel_filename({"outcome": "accepted"}) == "nfe-extraction.xlsx"


def test_the_workbook_follows_the_language_and_keeps_the_file_name() -> None:
    body = _payload(
        _extraction(
            line_items=(
                _line_item(
                    ncm="09012100",
                    cfop="0102",
                    quantity=Decimal(9),
                    unit_price=Decimal("8.00"),
                ),
            ),
            series="001",
            number="000000123",
            unsupported_components={
                "import_tax": True,
                "returned_ipi": False,
                "icms_relief": False,
                "fcp_st": False,
            },
        )
    )
    english = load_workbook(io.BytesIO(workbook_bytes(body, language="en")))
    portuguese = load_workbook(io.BytesIO(workbook_bytes(body, language="pt")))

    assert excel_filename(body) == "nfe-000000123-001.xlsx"
    assert portuguese.sheetnames == ["Resumo", "Cabeçalho", "Itens"]
    summary = portuguese["Resumo"]
    assert summary["A1"].value == "Resultado"
    assert summary["B1"].value == "Precisa de revisão"
    assert summary["A2"].value == "Mensagens da revisão"
    assert summary["B2"].value == (
        "Imposto de importação é um componente não suportado; "
        "a regra do valor total da nota fica de fora e o documento vai direto "
        "para revisão humana."
    )
    header = portuguese["Cabeçalho"]
    assert [cell.value for cell in header[1]] == [
        "Campo",
        "Valor",
        "Situação",
        "Mensagens",
    ]
    series = _header_row(header, "Série")
    assert series[1].value == "001"
    assert series[1].value == _header_row(english["Header"], "Series")[1].value
    items = portuguese["Itens"]
    assert [cell.value for cell in items[1]] == [
        "Código do produto",
        "Descrição",
        "NCM",
        "CFOP",
        "Quantidade",
        "Valor unitário",
        "Valor do item",
    ]
    assert _item_row(items, "SKU-1")[1].value == "Plain bearing"
    assert (
        _item_row(items, "SKU-1")[1].value
        == _item_row(english["Line items"], "SKU-1")[1].value
    )


def _gallery_body(caption: str) -> dict[str, object]:
    for item in load_gallery(_SYNTHETIC, today=_TODAY):
        if item.caption == caption or item.caption.startswith(caption + " ·"):
            return item.payload
    raise AssertionError(caption)


def _header_row(sheet: Worksheet, field: str) -> tuple:
    for row in sheet.iter_rows(min_row=2, max_col=4):
        if row[0].value == field:
            return row
    raise AssertionError(field)


def _item_row(sheet: Worksheet, product_code: str) -> tuple:
    for row in sheet.iter_rows(min_row=2, max_col=7):
        if row[0].value == product_code:
            return row
    raise AssertionError(product_code)


def _without_header_value(body: dict[str, object], name: str) -> dict[str, object]:
    extraction = body["extraction"]
    assert isinstance(extraction, dict)
    header = extraction["header"]
    assert isinstance(header, dict)
    return {
        **body,
        "extraction": {**extraction, "header": {**header, name: ""}},
    }


def _payload(extraction: Extraction) -> dict[str, object]:
    report = validate(extraction, today=_TODAY)
    outcome = (
        DocumentOutcome.ACCEPTED if report.accepted else DocumentOutcome.HUMAN_REVIEW
    )
    return document_payload(extraction, report, outcome=outcome, thread_id=None)


def _highlighted(cell: Cell) -> bool:
    if cell.fill.patternType != "solid":
        return False
    color = cell.fill.fgColor.rgb
    return isinstance(color, str) and color.endswith("FFF2CC")
